import io
import pytest
import docx
from reportlab.pdfgen import canvas
from fastapi.testclient import TestClient
from api.server import app

@pytest.fixture
def client():
    return TestClient(app)

def test_upload_txt(client):
    txt_data = io.BytesIO(b'Este es un reporte de prueba en texto plano con datos clave.')
    resp = client.post('/api/upload', files={'file': ('reporte.txt', txt_data, 'text/plain')})
    assert resp.status_code == 200
    res_body = resp.json()
    assert res_body['ok'] is True
    data = res_body['data']
    assert data['success'] is True
    assert data['file_type'] == 'text'
    assert 'reporte de prueba' in data['text']
    assert data['char_count'] > 20

def test_upload_docx(client):
    d = docx.Document()
    d.add_heading('Reporte Trimestral DOCX', level=1)
    d.add_paragraph('Este es el contenido de prueba para extraccion en Word.')
    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)

    resp = client.post('/api/upload', files={'file': ('reporte.docx', buf, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert resp.status_code == 200
    res_body = resp.json()
    assert res_body['ok'] is True
    data = res_body['data']
    assert data['success'] is True
    assert data['file_type'] == 'docx'
    assert 'Reporte Trimestral DOCX' in data['text']

def test_upload_pdf(client):
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 750, 'Reporte Oficial en Formato PDF')
    c.drawString(100, 730, 'Contenido relevante para JARVIS')
    c.save()
    buf.seek(0)

    resp = client.post('/api/upload', files={'file': ('reporte.pdf', buf, 'application/pdf')})
    assert resp.status_code == 200
    res_body = resp.json()
    assert res_body['ok'] is True
    data = res_body['data']
    assert data['success'] is True
    assert data['file_type'] == 'pdf'
    assert 'Reporte Oficial' in data['text']

def test_upload_invalid_extension(client):
    bin_data = io.BytesIO(b'\x00\x01\x02\x03\x04')
    resp = client.post('/api/upload', files={'file': ('binario.exe', bin_data, 'application/octet-stream')})
    assert resp.status_code == 400
    res_body = resp.json()
    assert res_body['ok'] is False
    assert res_body['error']['code'] == 'UNSUPPORTED_FORMAT'
