import os
import time
import glob
import json
import uuid
import asyncio
import urllib.parse
from contextlib import asynccontextmanager
from typing import Dict, Any, Optional, List
from fastapi import FastAPI, Request, File, UploadFile, WebSocket, WebSocketDisconnect, Query
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse, StreamingResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from core.config import settings, mask_key
from core.logger import logger
from core.orchestrator import JarvisOrchestrator
from core.memory_manager import memory_manager
from core.music_engine import music_engine
from core.notification_engine import ReminderScheduler, NotificationEngine
from core.voice_state_machine import VoiceStateMachine, VoiceState
from core.voice_pipeline import VoiceMetrics
from core.response_sanitizer import ResponseSanitizer
from core.document_generator import DocumentGenerator
from core.security import SecurityPolicy

# Proveedores directos para health checks
from providers.gemini_provider import GeminiProvider
from providers.openrouter_provider import OpenRouterProvider
from providers.anthropic_provider import AnthropicProvider

# Conexiones WebSocket activas para difusión de alarmas en vivo (BUG-31)
active_websockets: List[WebSocket] = []

@asynccontextmanager
async def lifespan(app: FastAPI):
    def _broadcast_alarm_to_clients(alarm_data: Dict[str, Any]):
        loop = asyncio.get_event_loop()
        for ws in list(active_websockets):
            async def _send(ws=ws, data=alarm_data):
                try:
                    await ws.send_json({"type": "reminder_alarm", **data})
                except Exception:
                    pass
            if loop.is_running():
                asyncio.run_coroutine_threadsafe(_send(), loop)
            else:
                loop.create_task(_send())

    NotificationEngine.register_broadcast_callback(_broadcast_alarm_to_clients)

    # Iniciar motor de recordatorios con APScheduler persistido en SQLite
    from core.reminders import reminder_engine
    reminder_engine.start()

    scheduler_task = asyncio.create_task(ReminderScheduler.start(interval_seconds=5))
    logger.info(f"[Servidor v{settings.VERSION}] ReminderScheduler y ReminderEngine iniciados en segundo plano.")
    
    # Auto-diagnóstico asíncrono no bloqueante en segundo plano (arranque instantáneo)
    async def _async_startup_health():
        try:
            health_report = await perform_health_check()
            logger.info(f"[Auto-Diagnóstico v3.0] Estado de Servicios:\n{json.dumps(health_report, indent=2)}")
        except Exception as e:
            logger.warning(f"[Auto-Diagnóstico] Error en chequeo inicial: {e}")

    asyncio.create_task(_async_startup_health())

    yield
    reminder_engine.stop()
    ReminderScheduler.stop()
    scheduler_task.cancel()
    logger.info("[Servidor] Motores de recordatorios detenidos.")

app = FastAPI(title="JARVIS AI Assistant", version=settings.VERSION, lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.ALLOWED_ORIGINS if settings.ALLOWED_ORIGINS != ["*"] else ["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

os.makedirs(settings.STATIC_DIR, exist_ok=True)
os.makedirs(settings.AUDIO_DIR, exist_ok=True)
app.mount("/static", StaticFiles(directory=str(settings.STATIC_DIR)), name="static")

orchestrator = JarvisOrchestrator()

def api_envelope(ok: bool, data: Any = None, error: Optional[Dict[str, Any]] = None, status_code: int = 200) -> JSONResponse:
    """Envelope unificado para APIs JSON según especificación JARVIS v4.0 Mobile-Ready (A.1)."""
    return JSONResponse(
        status_code=status_code,
        content={
            "ok": ok,
            "data": data,
            "error": error
        }
    )

class SessionStore:
    """
    Almacén de sesiones persistentes en memoria para clientes desktop y mobile (A.1).
    Permite tolerar reconexiones sin duplicar estado de sesión ni perder contexto conversacional.
    """
    def __init__(self):
        self._sessions: Dict[str, List[Dict[str, Any]]] = {}
        self._device_map: Dict[str, str] = {}
        self._last_seen: Dict[str, float] = {}

    def get_or_create(self, session_id: Optional[str] = None, device_id: Optional[str] = None) -> tuple[str, List[Dict[str, Any]]]:
        if not session_id and device_id and device_id in self._device_map:
            session_id = self._device_map[device_id]
        if not session_id:
            session_id = str(uuid.uuid4())
        if session_id not in self._sessions:
            self._sessions[session_id] = []
        if device_id:
            self._device_map[device_id] = session_id
        self._last_seen[session_id] = time.time()
        return session_id, self._sessions[session_id]

    def get_history(self, session_id: str) -> List[Dict[str, Any]]:
        return self._sessions.get(session_id, [])

    def save_history(self, session_id: str, history: List[Dict[str, Any]]):
        self._sessions[session_id] = history
        self._last_seen[session_id] = time.time()

    def reset(self, session_id: str):
        if session_id in self._sessions:
            self._sessions[session_id] = []

session_store = SessionStore()
connection_sessions: Dict[str, List[Dict[str, Any]]] = session_store._sessions

async def perform_health_check() -> Dict[str, Any]:
    gemini_p = GeminiProvider()
    openrouter_p = OpenRouterProvider()
    anthropic_p = AnthropicProvider()

    gemini_res, openrouter_res, anthropic_res, yt_res, spot_res = await asyncio.gather(
        gemini_p.health_check(),
        openrouter_p.health_check(),
        anthropic_p.health_check(),
        music_engine.youtube_provider.health_check(),
        music_engine.spotify_provider.health_check(),
        return_exceptions=True
    )

    def _fmt(r, name):
        if isinstance(r, Exception):
            return {"status": "error", "error": str(r)}
        return r

    return {
        "version": settings.VERSION,
        "gemini": _fmt(gemini_res, "gemini"),
        "openrouter": _fmt(openrouter_res, "openrouter"),
        "anthropic": _fmt(anthropic_res, "anthropic"),
        "youtube": _fmt(yt_res, "youtube"),
        "spotify": _fmt(spot_res, "spotify"),
        "timestamp": time.time()
    }

def cleanup_old_audio(max_age_seconds: int = 1800):
    try:
        now = time.time()
        for f in glob.glob(os.path.join(settings.AUDIO_DIR, "response_*.mp3")):
            if os.path.isfile(f) and (now - os.path.getmtime(f)) > max_age_seconds:
                try:
                    os.remove(f)
                except Exception:
                    pass
    except Exception:
        pass

@app.get("/", response_class=HTMLResponse)
async def serve_index():
    index_file = settings.TEMPLATES_DIR / "index.html"
    return FileResponse(index_file)

@app.get("/api/health")
async def health_endpoint():
    """Endpoint de auto-diagnóstico en tiempo real normalizado con envelope v4.0."""
    report = await perform_health_check()
    return api_envelope(ok=True, data=report)

_TTS_CACHE: Dict[str, bytes] = {}
_TTS_CACHE_MAX_SIZE = 300

@app.get("/api/tts/stream")
async def tts_stream_endpoint(text: str = Query(..., description="Texto a sintetizar con Edge-TTS")):
    clean_text = ResponseSanitizer.sanitize_for_speech(text)
    if not clean_text:
        return api_envelope(ok=False, data=None, error={"code": "BAD_REQUEST", "message": "Texto vacío"}, status_code=400)

    # Respuesta instantánea desde caché en memoria si ya fue sintetizada
    if clean_text in _TTS_CACHE:
        return Response(content=_TTS_CACHE[clean_text], media_type="audio/mpeg")

    async def _audio_generator():
        audio_buffer = bytearray() if len(clean_text) <= 350 else None
        try:
            import edge_tts
            communicate = edge_tts.Communicate(clean_text, voice="es-MX-JorgeNeural", rate="+16%")
            async for chunk in communicate.stream():
                if chunk["type"] == "audio":
                    data = chunk["data"]
                    if audio_buffer is not None:
                        audio_buffer.extend(data)
                    yield data
            if audio_buffer is not None and len(audio_buffer) > 0:
                if len(_TTS_CACHE) >= _TTS_CACHE_MAX_SIZE:
                    try:
                        oldest_key = next(iter(_TTS_CACHE))
                        del _TTS_CACHE[oldest_key]
                    except Exception:
                        pass
                _TTS_CACHE[clean_text] = bytes(audio_buffer)
        except Exception as e:
            logger.error(f"[EdgeTTS] Error en streaming: {e}")

    return StreamingResponse(_audio_generator(), media_type="audio/mpeg")

class ChatRequest(BaseModel):
    prompt: Optional[str] = None
    text: Optional[str] = None
    session_id: Optional[str] = None
    device_id: Optional[str] = None
    image: Optional[Dict[str, Any]] = None

class CreateReminderRequest(BaseModel):
    title: Optional[str] = None
    remind_at: Optional[str] = None
    priority: Optional[str] = "NORMAL"
    description: Optional[str] = ""
    # Aliases
    text: Optional[str] = None
    time: Optional[str] = None
    when: Optional[str] = None
    date: Optional[str] = None
    reminder: Optional[str] = None

class SnoozeReminderRequest(BaseModel):
    minutes: Optional[int] = 10


@app.post("/api/chat")
async def chat_endpoint(payload: ChatRequest):
    cleanup_old_audio()
    user_prompt = (payload.prompt or payload.text or "").strip()
    user_image = payload.image
    if not user_prompt and not user_image:
        return api_envelope(
            ok=False,
            data=None,
            error={"code": "BAD_REQUEST", "message": "El mensaje no puede estar vacío"},
            status_code=400
        )

    sid, history = session_store.get_or_create(session_id=payload.session_id, device_id=payload.device_id)
    result = await orchestrator.process_user_input(user_prompt, conversation_history=history, image_data=user_image)
    spoken_text = ResponseSanitizer.sanitize(result.get("response_text", ""))

    history.append({"role": "user", "content": user_prompt})
    history.append({"role": "assistant", "content": spoken_text})
    if len(history) > 16:
        history = history[-16:]
    session_store.save_history(sid, history)

    # URL-encode query string (BUG-08)
    encoded_text = urllib.parse.quote(spoken_text)
    audio_url = f"/api/tts/stream?text={encoded_text}"

    data_payload = {
        "text": spoken_text,
        "audio_url": audio_url,
        "playlist": result.get("playlist"),
        "current_track": result.get("current_track"),
        "tracks": result.get("tracks", []),
        "action_type": result.get("action_type"),
        "action_detail": result.get("action_detail"),
        "generated_image": result.get("generated_image"),
        "generated_doc": result.get("generated_doc"),
        "tools_executed": result.get("tools_executed", []),
        "provider": result.get("provider"),
        "session_id": sid,
        "device_id": payload.device_id
    }
    return api_envelope(ok=True, data=data_payload)

@app.post("/api/v1/message")
async def message_v1_endpoint(payload: ChatRequest):
    return await chat_endpoint(payload)

@app.post("/api/upload")
async def upload_file_endpoint(file: UploadFile = File(...)):
    """
    Recibe documentos (PDF, DOCX, TXT, CSV, MD, JSON) o imágenes.
    Extrae el texto de documentos y lo entrega listo para análisis conversacional por JARVIS.
    """
    if not file or not file.filename:
        return JSONResponse(status_code=400, content={"success": False, "error": "No se proporcionó ningún archivo."})

    filename = file.filename
    ext = os.path.splitext(filename)[1].lower()
    content_bytes = await file.read()
    file_size = len(content_bytes)

    if file_size > 25 * 1024 * 1024:
        return JSONResponse(status_code=400, content={"success": False, "error": "El archivo excede el límite máximo de 25MB."})

    extracted_text = ""
    file_type = "unknown"
    data_url = None

    try:
        # 1. Documentos PDF
        if ext == ".pdf":
            file_type = "pdf"
            import io
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content_bytes))
            pages_text = []
            max_pages = min(len(reader.pages), 50)
            for p_num in range(max_pages):
                p_text = reader.pages[p_num].extract_text() or ""
                if p_text.strip():
                    pages_text.append(f"--- Página {p_num + 1} ---\n{p_text.strip()}")
            extracted_text = "\n\n".join(pages_text)

        # 2. Documentos DOCX / Word
        elif ext in [".docx", ".doc"]:
            file_type = "docx"
            import io
            import docx
            doc = docx.Document(io.BytesIO(content_bytes))
            paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_txt:
                        paras.append(f"| {row_txt} |")
            extracted_text = "\n\n".join(paras)

        # 3. Archivos de texto plano / Markdown / CSV / JSON
        elif ext in [".txt", ".md", ".csv", ".json", ".log", ".py", ".js", ".html"]:
            file_type = "text"
            try:
                extracted_text = content_bytes.decode("utf-8")
            except UnicodeDecodeError:
                extracted_text = content_bytes.decode("latin-1", errors="replace")

        # 4. Imágenes
        elif ext in [".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"]:
            file_type = "image"
            import base64
            mime = file.content_type or "image/png"
            b64 = base64.b64encode(content_bytes).decode("utf-8")
            data_url = f"data:{mime};base64,{b64}"
            try:
                pics_dir = SecurityPolicy.get_actual_user_dir("Pictures")
                pics_dir.mkdir(parents=True, exist_ok=True)
                saved_path = pics_dir / filename
                with open(saved_path, "wb") as f:
                    f.write(content_bytes)
                DocumentGenerator.set_last_session_image(saved_path)
                logger.info(f"[FileUpload] Imagen guardada en disco: {saved_path}")
            except Exception as se:
                logger.warning(f"[FileUpload] No se pudo guardar copia de imagen en disco: {se}")

        else:
            return api_envelope(
                ok=False,
                data=None,
                error={
                    "code": "UNSUPPORTED_FORMAT",
                    "message": f"Formato no soportado ({ext}). Formatos válidos: PDF, DOCX, TXT, CSV, MD, imágenes."
                },
                status_code=400
            )

        logger.info(f"[FileUpload] Archivo recibido: '{filename}' ({file_type}, {file_size} bytes, texto extraído: {len(extracted_text)} chars)")
        upload_data = {
            "success": True,
            "filename": filename,
            "file_type": file_type,
            "size_bytes": file_size,
            "char_count": len(extracted_text),
            "text": extracted_text[:30000],
            "data_url": data_url,
            "message": f"Archivo '{filename}' procesado correctamente."
        }
        return api_envelope(ok=True, data=upload_data)

    except Exception as e:
        logger.error(f"[FileUpload] Error procesando archivo '{filename}': {e}")
        return api_envelope(ok=False, data=None, error={"code": "INTERNAL_ERROR", "message": str(e)}, status_code=500)

@app.get("/api/profile")
async def get_profile_endpoint():
    return api_envelope(ok=True, data=memory_manager.get_all_profile())

@app.get("/api/reminders")
async def get_reminders_endpoint(include_completed: bool = False):
    return api_envelope(ok=True, data=memory_manager.list_reminders(include_completed=include_completed))

@app.post("/api/reminders")
async def create_reminder_endpoint(payload: CreateReminderRequest):
    clean_title = (
        payload.title or payload.text or payload.reminder or "Recordatorio"
    ).strip()
    clean_remind_at = (
        payload.remind_at or payload.time or payload.when or payload.date or payload.description or "en 1 hora"
    ).strip()
    clean_priority = (payload.priority or "NORMAL").upper()

    res = memory_manager.create_reminder(
        title=clean_title,
        remind_at_expression=clean_remind_at,
        priority=clean_priority,
        description=payload.description or ""
    )
    if res.get("success"):
        return api_envelope(ok=True, data=res.get("data"))
    return api_envelope(ok=False, data=None, error={"code": "CREATION_FAILED", "message": res.get("error", "Error creando recordatorio")}, status_code=400)

@app.delete("/api/reminders/{reminder_id}")
async def delete_reminder_endpoint(reminder_id: int):
    res = memory_manager.delete_reminder(reminder_id)
    if res.get("success"):
        return api_envelope(ok=True, data=res)
    return api_envelope(ok=False, data=None, error={"code": "NOT_FOUND", "message": res.get("error", "No encontrado")}, status_code=404)

@app.post("/api/reminders/{reminder_id}/snooze")
async def snooze_reminder_endpoint(reminder_id: int, payload: Optional[SnoozeReminderRequest] = None):
    mins = payload.minutes if payload and payload.minutes else 10
    res = memory_manager.snooze_reminder(reminder_id=reminder_id, minutes=mins)
    if res.get("success"):
        return api_envelope(ok=True, data=res.get("data"))
    return api_envelope(ok=False, data=None, error={"code": "SNOOZE_FAILED", "message": res.get("error", "Error al posponer")}, status_code=400)

@app.post("/api/reminders/{reminder_id}/complete")
async def complete_reminder_endpoint(reminder_id: int):
    ok = memory_manager.complete_reminder(reminder_id=reminder_id)
    if ok:
        return api_envelope(ok=True, data={"id": reminder_id, "status": "COMPLETED"})
    return api_envelope(ok=False, data=None, error={"code": "NOT_FOUND", "message": f"No se encontró recordatorio #{reminder_id}"}, status_code=404)

@app.delete("/api/reminders")
async def clear_reminders_endpoint():
    res = memory_manager.clear_all_reminders()
    return api_envelope(ok=True, data=res)

@app.get("/api/pending")
async def get_pending_endpoint():
    return api_envelope(ok=True, data=memory_manager.get_pending_summary())

@app.post("/api/reset")
async def reset_endpoint(session_id: Optional[str] = None):
    sid = session_id or "http_default"
    session_store.reset(sid)
    return api_envelope(ok=True, data={"status": "ok", "message": "Conversación reiniciada para la sesión.", "session_id": sid})

@app.get("/api/chat/stream")
@app.get("/api/v1/chat/stream")
async def chat_sse_get_endpoint(
    prompt: str = Query(...),
    session_id: Optional[str] = None,
    device_id: Optional[str] = None
):
    """
    Streaming mediante Server-Sent Events (SSE) token por token y oración por oración.
    Permite a clientes web y móviles recibir texto en tiempo real con latencia mínima y URLs de audio TTS por oración.
    """
    user_prompt = prompt.strip()
    if not user_prompt:
        return api_envelope(ok=False, data=None, error={"code": "BAD_REQUEST", "message": "El mensaje no puede estar vacío"}, status_code=400)

    sid, history = session_store.get_or_create(session_id=session_id, device_id=device_id)

    async def sse_generator():
        try:
            async for chunk in orchestrator.process_user_input_stream(user_prompt, conversation_history=history):
                c_type = chunk.get("type")
                if c_type == "token":
                    data = json.dumps({"type": "token", "token": chunk.get("content", "")})
                    yield f"event: token\ndata: {data}\n\n"
                elif c_type == "sentence":
                    clean = ResponseSanitizer.sanitize(chunk["text"])
                    spoken = ResponseSanitizer.sanitize_for_speech(clean)
                    encoded = urllib.parse.quote(spoken)
                    data = json.dumps({
                        "type": "sentence",
                        "text": clean,
                        "audio_url": f"/api/tts/stream?text={encoded}",
                        "is_final": chunk.get("is_final", False)
                    })
                    yield f"event: sentence\ndata: {data}\n\n"
                elif c_type == "tool_executed":
                    data = json.dumps({
                        "type": "tool_executed",
                        "tool_name": chunk.get("tool_name"),
                        "result": chunk.get("result")
                    })
                    yield f"event: tool_executed\ndata: {data}\n\n"
                elif c_type == "response_end":
                    final_t = chunk.get("full_text", "")
                    history.append({"role": "user", "content": user_prompt})
                    history.append({"role": "assistant", "content": final_t})
                    if len(history) > 16:
                        history[:] = history[-16:]
                    session_store.save_history(sid, history)
                    data = json.dumps({
                        "type": "response_end",
                        "full_text": final_t,
                        "generated_image": chunk.get("generated_image"),
                        "generated_doc": chunk.get("generated_doc")
                    })
                    yield f"event: response_end\ndata: {data}\n\n"
        except Exception as ex:
            logger.error(f"[SSE Stream] Error: {ex}")
            yield f"event: error\ndata: {json.dumps({'type': 'error', 'message': str(ex)})}\n\n"

    return StreamingResponse(
        sse_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no"
        }
    )

@app.post("/api/chat/stream")
@app.post("/api/v1/chat/stream")
async def chat_sse_post_endpoint(payload: ChatRequest):
    return await chat_sse_get_endpoint(
        prompt=payload.prompt or payload.text or "",
        session_id=payload.session_id,
        device_id=payload.device_id
    )

@app.websocket("/ws/v1/chat")
async def websocket_chat_endpoint(
    websocket: WebSocket,
    session_id: Optional[str] = Query(None),
    device_id: Optional[str] = Query(None)
):
    await websocket.accept()
    resolved_sid, history = session_store.get_or_create(session_id=session_id, device_id=device_id)
    active_websockets.append(websocket)
    logger.info(f"[WebSocket] Cliente conectado con sesión: {resolved_sid} (device: {device_id})")

    # Enviar sincronización de estado inicial (BUG-28, BUG-32 & Perfil Completo)
    try:
        pending_data = memory_manager.get_pending_summary()
        profile = memory_manager.get_all_profile()
        user_name = profile.get("full_name") or profile.get("name") or "Usuario"
        alias = profile.get("alias")
        active_provider = orchestrator.primary_provider.name if hasattr(orchestrator, "primary_provider") else "OpenRouter"

        await websocket.send_json({
            "type": "sync_state",
            "session_id": resolved_sid,
            "device_id": device_id,
            "user_name": user_name,
            "alias": alias,
            "user_profile": profile,
            "active_reminders_count": len(pending_data.get("reminders", [])),
            "active_tasks_count": len(pending_data.get("tasks", [])),
            "reminders": pending_data.get("reminders", [])[:5],
            "tasks": pending_data.get("tasks", [])[:5],
            "provider_info": active_provider
        })
    except Exception as sync_err:
        logger.warning(f"[WebSocket] Error enviando sync_state: {sync_err}")

    state_machine = VoiceStateMachine(
        on_state_change=lambda old, new, gid: asyncio.create_task(
            websocket.send_json({
                "type": "state_change",
                "state": new.value,
                "generation_id": gid
            })
        )
    )

    current_task: Optional[asyncio.Task] = None

    try:
        while True:
            data = await websocket.receive_text()
            try:
                msg = json.loads(data)
            except json.JSONDecodeError:
                await websocket.send_json({"type": "error", "message": "JSON inválido"})
                continue

            msg_type = msg.get("type")

            if msg_type == "ping":
                await websocket.send_json({
                    "type": "pong",
                    "session_id": resolved_sid,
                    "timestamp": time.time()
                })
                continue

            if msg_type == "barge_in":
                logger.info(f"[WebSocket] Barge-in recibido para sesión {resolved_sid}")
                state_machine.trigger_barge_in()
                if current_task and not current_task.done():
                    current_task.cancel()
                continue

            if msg_type == "user_text":
                user_prompt = msg.get("text", "").strip()
                user_image = msg.get("image")
                if not user_prompt and not user_image:
                    continue

                if current_task and not current_task.done():
                    current_task.cancel()

                gen_id = state_machine.start_new_generation()
                metrics = VoiceMetrics()
                metrics.mark_speech_end()
                metrics.mark_stt_final()

                state_machine.transition_to(VoiceState.PROCESSING, reason="Procesando input")

                async def _stream_handler(active_gid: int, sid: str):
                    accumulated_full_text = ""
                    hist = session_store.get_history(sid)
                    first_token_emitted = False
                    first_sentence_emitted = False
                    try:
                        async for chunk in orchestrator.process_user_input_stream(user_prompt, hist, image_data=user_image):
                            if not state_machine.is_generation_valid(active_gid):
                                logger.info(f"[WebSocket] Descartando chunk gen #{active_gid} (antigua)")
                                break

                            if chunk.get("type") == "token":
                                token_txt = chunk.get("content", "")
                                if not first_token_emitted:
                                    first_token_emitted = True
                                    metrics.mark_llm_first_token()
                                await websocket.send_json({
                                    "type": "token",
                                    "token": token_txt,
                                    "generation_id": active_gid
                                })

                            elif chunk.get("type") == "sentence":
                                if not first_sentence_emitted:
                                    first_sentence_emitted = True
                                    metrics.mark_sentence_ready()
                                    metrics.mark_tts_start()
                                    state_machine.transition_to(VoiceState.SPEAKING, reason="Iniciando TTS")

                                clean_chunk = ResponseSanitizer.sanitize(chunk["text"])
                                if clean_chunk:
                                    accumulated_full_text += ("\n" if accumulated_full_text and "\n" in clean_chunk else " ") + clean_chunk
                                    spoken_chunk = ResponseSanitizer.sanitize_for_speech(clean_chunk)
                                    encoded_chunk = urllib.parse.quote(spoken_chunk)
                                    await websocket.send_json({
                                        "type": "tts_chunk",
                                        "text": clean_chunk,
                                        "audio_url": f"/api/tts/stream?text={encoded_chunk}",
                                        "generation_id": active_gid,
                                        "is_final": chunk.get("is_final", False)
                                    })

                            elif chunk.get("type") == "tool_executed":
                                t_name = chunk["tool_name"]
                                t_result = chunk["result"]

                                await websocket.send_json({
                                    "type": "tool_executed",
                                    "tool_name": t_name,
                                    "result": t_result,
                                    "generation_id": active_gid
                                })

                                # Control bidireccional de reproductor multimedia (BUG-06)
                                if t_name == "pause_music":
                                    await websocket.send_json({
                                        "type": "player_control",
                                        "action": "pause",
                                        "generation_id": active_gid
                                    })
                                elif t_name == "resume_music":
                                    await websocket.send_json({
                                        "type": "player_control",
                                        "action": "play",
                                        "generation_id": active_gid
                                    })
                                elif t_name == "skip_music":
                                    await websocket.send_json({
                                        "type": "player_control",
                                        "action": "next",
                                        "data": t_result.get("data", {}),
                                        "generation_id": active_gid
                                    })
                                elif t_name in ["play_music", "recommend_music"] and t_result.get("success"):
                                    d = t_result.get("data", {})
                                    await websocket.send_json({
                                        "type": "player_control",
                                        "action": "play_track",
                                        "data": {
                                             "current_track": d.get("current_track"),
                                            "tracks": d.get("tracks", []),
                                            "playlist": d.get("playlist", [])
                                        },
                                        "generation_id": active_gid
                                    })
                                elif t_name == "clear_music_queue" and t_result.get("success"):
                                    await websocket.send_json({
                                        "type": "player_control",
                                        "action": "clear_queue",
                                        "generation_id": active_gid
                                    })

                            elif chunk.get("type") == "response_end":
                                metrics.mark_response_end()
                                rep = metrics.calculate_report()
                                final_t = ResponseSanitizer.sanitize(chunk.get("full_text") or accumulated_full_text)
                                
                                # Registrar en historial aislado de la sesión (BUG-11, A.1)
                                hist.append({"role": "user", "content": user_prompt})
                                hist.append({"role": "assistant", "content": final_t})
                                if len(hist) > 16:
                                    hist[:] = hist[-16:]
                                session_store.save_history(sid, hist)

                                spoken_final = ResponseSanitizer.sanitize_for_speech(final_t)
                                encoded_final = urllib.parse.quote(spoken_final)
                                await websocket.send_json({
                                    "type": "response_end",
                                    "full_text": final_t,
                                    "audio_url": f"/api/tts/stream?text={encoded_final}",
                                    "generated_image": chunk.get("generated_image"),
                                    "generated_doc": chunk.get("generated_doc"),
                                    "generation_id": active_gid,
                                    "metrics": rep
                                })
                                state_machine.transition_to(VoiceState.IDLE, reason="Fin de respuesta")

                    except asyncio.CancelledError:
                        logger.info(f"[WebSocket] Tarea gen #{active_gid} cancelada por Barge-in.")
                    except Exception as ex:
                        logger.error(f"[WebSocket] Error en stream: {ex}")
                        state_machine.transition_to(VoiceState.ERROR, reason=str(ex))
                        await websocket.send_json({"type": "error", "message": str(ex)})

                current_task = asyncio.create_task(_stream_handler(gen_id, resolved_sid))

    except WebSocketDisconnect:
        logger.info(f"[WebSocket] Cliente desconectado ({resolved_sid}). Estado conservado para reconexión móvil.")
        if current_task and not current_task.done():
            current_task.cancel()
    finally:
        if websocket in active_websockets:
            active_websockets.remove(websocket)
